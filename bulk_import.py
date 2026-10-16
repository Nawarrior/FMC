"""Parse uploads, validate cards in batches and atomically confirm owned imports."""
import csv
from datetime import datetime
from io import BytesIO, StringIO
import json
from pathlib import Path
import re
import unicodedata
from urllib.error import HTTPError, URLError
from urllib.request import Request, urlopen
from uuid import uuid4
import zipfile

import database

MAX_ROWS = 500
MAX_BYTES = 2 * 1024 * 1024
CHUNK_SIZE = 25


def normalize(value) -> str:
    return "".join(c for c in unicodedata.normalize("NFKD", str(value or "").strip().lower()) if not unicodedata.combining(c))


HEADERS = {
    "nombre": "name", "carta": "name", "name": "name", "card": "name", "card name": "name",
    "cantidad": "quantity", "quantity": "quantity", "qty": "quantity",
    "edicion": "set", "set": "set", "codigo": "set", "set code": "set",
    "numero": "number", "number": "number", "collector_number": "number",
    "nota": "note", "note": "note", "estado": "intent", "intent": "intent",
}


def parse_rows(matrix, default_intent: str) -> list[dict]:
    if default_intent not in {"busco", "ofrezco"}:
        raise ValueError("Selecciona Busco u Ofrezco.")
    matrix = [(index, row) for index, row in enumerate(matrix, 1) if any(value is not None and str(value).strip() for value in row)]
    if not matrix:
        raise ValueError("El listado esta vacio.")
    first = [HEADERS.get(normalize(value)) for value in matrix[0][1]]
    has_header = "name" in first
    columns = first if has_header else ["name", "quantity", "set", "number", "note", "intent"]
    if has_header:
        matrix = matrix[1:]
    if not matrix or len(matrix) > MAX_ROWS:
        raise ValueError(f"Debes incluir entre 1 y {MAX_ROWS} filas de cartas.")
    result = []
    for line, values in matrix:
        row = {key: values[index] for index, key in enumerate(columns) if key and index < len(values)}
        entry = {"line": line, "name": str(row.get("name") or "").strip(), "quantity": 1,
                 "set": str(row.get("set") or "").strip().lower(),
                 "number": str(row.get("number") or "").strip(), "note": str(row.get("note") or "").strip(),
                 "intent": normalize(row.get("intent")) or default_intent, "status": "pending", "error": ""}
        try:
            if any(isinstance(value, str) and value.startswith("=") for value in values):
                raise ValueError("Reemplaza las formulas por valores antes de importar.")
            raw_quantity = row.get("quantity")
            if raw_quantity is None or str(raw_quantity).strip() == "":
                raw_quantity = 1
            number = float(raw_quantity)
            if not number.is_integer() or not 1 <= number <= 9999:
                raise ValueError("La cantidad debe ser un entero entre 1 y 9999.")
            entry["quantity"] = int(number)
            if not entry["name"] or len(entry["name"]) > 200:
                raise ValueError("Indica el nombre completo de la carta (maximo 200 caracteres).")
            if entry["intent"] not in {"busco", "ofrezco"}:
                raise ValueError("El estado debe ser Busco u Ofrezco.")
            if entry["set"] and not re.fullmatch(r"[a-z0-9]{2,8}", entry["set"]):
                raise ValueError("La edicion debe ser su codigo, por ejemplo M21 o DMU.")
            if entry["number"] and not entry["set"]:
                raise ValueError("Para indicar un numero de coleccion debes indicar la edicion.")
            if len(entry["note"]) > 500 or len(entry["number"]) > 20:
                raise ValueError("La nota o el numero de coleccion es demasiado largo.")
        except (TypeError, ValueError, OverflowError) as exc:
            entry.update(status="error", error=str(exc))
        result.append(entry)
    return result


def parse_upload(data: bytes, filename: str, default_intent: str) -> list[dict]:
    if len(data) > MAX_BYTES:
        raise ValueError("El archivo no debe superar 2 MB.")
    suffix = Path(filename).suffix.lower()
    if suffix == ".xlsx":
        from openpyxl import load_workbook
        try:
            with zipfile.ZipFile(BytesIO(data)) as archive:
                if len(archive.infolist()) > 2000 or sum(item.file_size for item in archive.infolist()) > 30 * 1024 * 1024:
                    raise ValueError("El Excel supera el tamano permitido al descomprimirlo.")
            workbook = load_workbook(BytesIO(data), read_only=True, data_only=False, keep_links=False)
            try:
                sheet = workbook.active
                if sheet is None or sheet.max_row and sheet.max_row > MAX_ROWS + 1:
                    raise ValueError(f"Usa la primera hoja con un maximo de {MAX_ROWS} cartas.")
                matrix = list(sheet.iter_rows(max_row=MAX_ROWS + 2, max_col=12, values_only=True))
            finally:
                workbook.close()
        except ValueError:
            raise
        except Exception as exc:
            raise ValueError("No se pudo leer el Excel. Usa un archivo .xlsx valido.") from exc
        return parse_rows(matrix, default_intent)
    if suffix not in {".txt", ".csv", ".tsv"}:
        raise ValueError("Usa Excel .xlsx, CSV o texto .txt. Convierte los archivos .xls a .xlsx.")
    try:
        text = data.decode("utf-16") if data.startswith((b"\xff\xfe", b"\xfe\xff")) else data.decode("utf-8-sig")
    except UnicodeDecodeError:
        text = data.decode("cp1252")
    lines = [line for line in text.splitlines() if line.strip() and not line.lstrip().startswith(("#", "//"))]
    if len(lines) > MAX_ROWS + 1:
        raise ValueError(f"El listado no puede superar {MAX_ROWS} filas de cartas.")
    if suffix in {".csv", ".tsv"} or lines and any(separator in lines[0] for separator in ("\t", ";")):
        try:
            dialect = csv.Sniffer().sniff("\n".join(lines[:5]), delimiters=",;\t")
            matrix = list(csv.reader(StringIO("\n".join(lines)), dialect))
        except csv.Error:
            matrix = [[line] for line in lines]
    else:
        matrix = []
        for line in lines:
            quantity = 1
            match = re.match(r"^\s*(\d+)\s*[xX]?\s+(.+?)\s*$", line)
            name = line.strip()
            if match:
                quantity, name = match.groups()
            edition = re.match(r"^(.*?)\s+\(([A-Za-z0-9]{2,8})\)(?:\s+(\S+))?$", name)
            matrix.append([edition[1] if edition else name, quantity, edition[2] if edition else "", edition[3] if edition else ""])
    return parse_rows(matrix, default_intent)


def create_import(owner_id: str, rows: list[dict]) -> str:
    batch_id = uuid4().hex
    columns = tuple(database.TABLE_DEFINITIONS["card_imports"]["columns"])
    item = {"import_id": batch_id, "user_id": owner_id, "rows": rows, "state": "draft", "created_at": datetime.now().isoformat()}
    with database.get_connection() as db:
        db.execute(f"INSERT INTO card_imports (position, {', '.join(columns)}) VALUES (0, {', '.join('?' for _ in columns)})",
                   tuple(database._serialize_value("card_imports", key, item[key]) for key in columns))
    return batch_id


def get_import(batch_id: str, owner_id: str) -> dict | None:
    with database.get_connection() as db:
        row = db.execute("SELECT * FROM card_imports WHERE import_id = ? AND user_id = ?", (batch_id, owner_id)).fetchone()
    if not row:
        return None
    return {**dict(row), "rows": json.loads(row["rows"])}


def fetch_collection(identifiers: list[dict]) -> list[dict]:
    req = Request("https://api.scryfall.com/cards/collection", data=json.dumps({"identifiers": identifiers}).encode(),
                  headers={"User-Agent": "FMC-MTG-Search/1.0", "Accept": "application/json", "Content-Type": "application/json"})
    try:
        with urlopen(req, timeout=15) as response:
            return json.loads(response.read().decode("utf-8")).get("data", [])
    except (HTTPError, URLError, TimeoutError, ValueError) as exc:
        raise RuntimeError("No se pudo validar con Scryfall. Puedes reintentar sin volver a subir el archivo.") from exc


def card_matches(row: dict, card: dict) -> bool:
    names = {normalize(card.get("name")), normalize(card.get("printed_name"))}
    names.update(normalize(face.get("name")) for face in card.get("card_faces", []))
    return (normalize(row["name"]) in names
            and (not row["set"] or row["set"] == card.get("set"))
            and (not row["number"] or row["number"] == card.get("collector_number")))


def resolve_import(batch_id: str, owner_id: str) -> dict:
    batch = get_import(batch_id, owner_id)
    if not batch or batch["state"] != "draft":
        raise ValueError("La importacion no esta disponible.")
    pending = [row for row in batch["rows"] if row["status"] == "pending"][:CHUNK_SIZE]
    identifiers = []
    for row in pending:
        identifier = {"name": row["name"]}
        if row["set"]:
            identifier["set"] = row["set"]
        if row["number"]:
            identifier = {"set": row["set"], "collector_number": row["number"]}
        if identifier not in identifiers:
            identifiers.append(identifier)
    if pending:
        cards = fetch_collection(identifiers)
        for row in pending:
            card = next((card for card in cards if card_matches(row, card)), None)
            if not card:
                row.update(status="error", error="No encontrada. Revisa el nombre, la edicion y el numero.")
            else:
                faces = card.get("card_faces") or [{}]
                image = card.get("image_uris") or faces[0].get("image_uris") or {}
                row.update(status="ready", card={"id": card["id"], "name": card["name"], "set": card.get("set", ""),
                           "set_name": card.get("set_name", ""), "rarity": card.get("rarity", ""), "image_url": image.get("normal", "")})
    with database.get_connection() as db:
        user = db.execute("SELECT email FROM users WHERE user_id = ?", (owner_id,)).fetchone()
        email = user[0].strip().lower() if user and user[0] else ""
        existing = {(row[0], row[1]) for row in db.execute("SELECT card_id, intent FROM saved_cards WHERE (user_id = ? OR lower(trim(user_email)) = ?) AND COALESCE(transaction_completed, 0) = 0", (owner_id, email)).fetchall()}
        for row in batch["rows"]:
            if row["status"] == "ready":
                key = (row["card"]["id"], row["intent"])
                if key in existing:
                    row.update(status="duplicate", error="Ya existe en tu listado o en una fila anterior de esta carga.")
                else:
                    existing.add(key)
        db.execute("UPDATE card_imports SET rows = ? WHERE import_id = ? AND user_id = ? AND state = 'draft'", (json.dumps(batch["rows"]), batch_id, owner_id))
    return batch


def confirm_import(batch_id: str, owner: dict, selected: set[int]) -> list[dict]:
    with database.get_connection() as db:
        db.execute("BEGIN IMMEDIATE")
        batch = db.execute("SELECT * FROM card_imports WHERE import_id = ? AND user_id = ?", (batch_id, owner["user_id"])).fetchone()
        if not batch or batch["state"] != "draft":
            raise ValueError("Esta carga ya fue importada o no te pertenece.")
        rows = json.loads(batch["rows"])
        if any(row["status"] == "pending" for row in rows):
            raise ValueError("Espera a que termine la validacion antes de importar.")
        if not selected or any(index < 0 or index >= len(rows) or rows[index]["status"] != "ready" for index in selected):
            raise ValueError("Selecciona al menos una carta validada para importar.")
        existing = {(row[0], row[1]) for row in db.execute("SELECT card_id, intent FROM saved_cards WHERE (user_id = ? OR lower(trim(user_email)) = ?) AND COALESCE(transaction_completed, 0) = 0", (owner["user_id"], owner["email"].strip().lower())).fetchall()}
        position = db.execute("SELECT COALESCE(MAX(position), -1) + 1 FROM saved_cards").fetchone()[0]
        added = []
        columns = tuple(database.TABLE_DEFINITIONS["saved_cards"]["columns"])
        for index in sorted(selected):
            row = rows[index]
            card = row["card"]
            key = (card["id"], row["intent"])
            if key in existing:
                continue
            existing.add(key)
            item = {"card_id": card["id"], "name": card["name"], "edition": card["set_name"], "edition_code": card["set"].upper(),
                    "rarity": card["rarity"], "image_url": card["image_url"], "quantity": row["quantity"], "intent": row["intent"], "note": row["note"],
                    "user_id": owner["user_id"], "user_name": owner["name"], "user_email": owner["email"],
                    "transaction_completed": False, "created_at": datetime.now().strftime("%Y-%m-%d %H:%M:%S")}
            db.execute(f"INSERT INTO saved_cards (position, {', '.join(columns)}) VALUES (?, {', '.join('?' for _ in columns)})",
                       (position + len(added), *(database._serialize_value("saved_cards", key, item.get(key)) for key in columns)))
            added.append(item)
        db.execute("UPDATE card_imports SET state = 'imported' WHERE import_id = ? AND user_id = ?", (batch_id, owner["user_id"]))
    return added
